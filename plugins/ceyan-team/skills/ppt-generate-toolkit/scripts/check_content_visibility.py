#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""内容可见性体检 —— 查「内容在不在画面里」，与 check_editability.py 互补。

## 为什么需要这个脚本

check_editability.py 只回答一个问题：**这一页是不是拿图片冒充的**。
它**完全不查**内容有没有落在画布外、文本框装不装得下、装不下有没有压住下面的东西。
两个判据正交，拿一个当两个用就会漏。

真实事故（某文旅项目实测）：check_editability 报全页合格、图片覆盖 0%，
而实际有 14 页内容落在画布外、多页文本排不下，其中一页丢了整段文字正文、
一页整张表在 y 远超画布（画布 1080 之外）。

## 它查三样

1. **越界** —— 形状右/下边缘超出画布（x<0 的出血装饰不算）
2. **溢出** —— 文本框按字体实测宽度换行后，所需高度超过框高
3. **碰撞** —— 溢出的文本压到下方同列的其它形状（垂直侵入 > 阈值）

## 退出码

    0 = 无问题        1 = 有越界或溢出        2 = 有碰撞（最严重）

用法：
    python check_content_visibility.py deck.pptx
    python check_content_visibility.py deck.pptx --json out.json
    python check_content_visibility.py deck.pptx --min-font 8 --tolerance 1.5
    python check_content_visibility.py deck.pptx --fonts "C:/Windows/Fonts/NotoSansSC-VF.ttf"

依赖：python-pptx、Pillow（缺 Pillow 时退化为按字符数估算，精度下降但能跑）
"""
from __future__ import annotations

import argparse
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.exit('[FATAL] 缺 python-pptx。装它：pip install python-pptx')

EMU_PER_PT = 12700

# ── 字体单行高系数 ────────────────────────────────────────────────────────────
# PowerPoint 的「倍数行距」乘的是**字体自带的单行高**，不是字号。
# Noto Sans SC 实测单行高 = 1.20em（PIL 实测 ascent 100 / descent 20 @100px）。
# 这条是踩过的坑：旧稿按 1.49em 估，实际渲染 1.86em，整整少算 25%，
# 于是「构建认为排得下、审计判为溢出」，自己跟自己打架。
# 教训：能用绝对行距（spcPts）就别用倍数；非用倍数不可时，系数必须实测。
SINGLE_LINE_RATIO = 1.20

# 各字体族的单行高系数（未列出的走 SINGLE_LINE_RATIO）
FONT_RATIO = {
    'noto sans sc': 1.20, 'noto serif sc': 1.24,
    'microsoft yahei': 1.32, '微软雅黑': 1.32,
    'simsun': 1.14, '宋体': 1.14,
    'simhei': 1.16, '黑体': 1.16,
    'arial': 1.15, 'helvetica': 1.15, 'calibri': 1.22,
}

DEFAULT_FONT_CANDIDATES = [
    r'C:\Windows\Fonts\NotoSansSC-VF.ttf',
    r'C:\Windows\Fonts\msyh.ttc',
    r'C:\Windows\Fonts\msyhbd.ttc',
    r'C:\Windows\Fonts\simsun.ttc',
    '/System/Library/Fonts/PingFang.ttc',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',
]

_font_cache = {}
_font_path = None
_have_pil = True


def _init_font(explicit: str | None) -> tuple[str | None, bool]:
    """确定量字用的字体文件。返回 (路径, 是否有 PIL)。"""
    try:
        from PIL import ImageFont  # noqa: F401
    except ImportError:
        return None, False
    for p in ([explicit] if explicit else []) + DEFAULT_FONT_CANDIDATES:
        if p and os.path.exists(p):
            return p, True
    return None, True


def _fnt(px: float):
    from PIL import ImageFont
    k = max(6, int(round(px)))
    if k not in _font_cache:
        _font_cache[k] = ImageFont.truetype(_font_path, k)
    return _font_cache[k]


def _text_width(text: str, px: float) -> float:
    """一行文字的实际像素宽度。无 PIL 时按东亚字 1.0em / 拉丁 0.55em 估。"""
    if _have_pil and _font_path:
        f = _fnt(px)
        return sum(f.getlength(ch) for ch in text)
    w = 0.0
    for ch in text:
        w += px if ord(ch) > 0x2E80 else px * 0.55
    return w


def _wrap_lines(text: str, px: float, box_w: float) -> int:
    """按实际字宽换行，返回行数。"""
    if not text:
        return 0
    total = 0
    for seg in text.split('\n'):
        if not seg:
            total += 1
            continue
        n, cur = 1, 0.0
        for ch in seg:
            w = _text_width(ch, px)
            if cur + w > box_w and cur > 0:
                n += 1
                cur = w
            else:
                cur += w
        total += n
    return total


def _px_per_pt(prs) -> float:
    """本稿 1pt 折多少 px —— 由画布宽度反推，不假设 96dpi。"""
    # 常见 16:9 稿按 1920 宽输出；若画布是 13.333in = 960pt，则 1pt = 2px
    w_pt = prs.slide_width / EMU_PER_PT
    target_px = 1920.0 if abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.02 else 1440.0
    return target_px / w_pt if w_pt else 2.0


def _iter_shapes(shapes):
    """递归展开组合，组合里的形状也要查。"""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _font_family(sh, para) -> str:
    for r in para.runs:
        if r.font.name:
            return r.font.name
    try:
        return sh.text_frame.paragraphs[0].runs[0].font.name or ''
    except Exception:
        return ''


def audit(pptx: str, min_font: float, tolerance: float, explicit_font: str | None):
    global _font_path, _have_pil
    _font_path, _have_pil = _init_font(explicit_font)

    prs = Presentation(pptx)
    k = _px_per_pt(prs)                      # pt → px
    CANVAS_W = prs.slide_width / EMU_PER_PT * k
    CANVAS_H = prs.slide_height / EMU_PER_PT * k
    E = EMU_PER_PT / k                       # 1 px 等于多少 EMU（k=2 时即 6350）

    pages = []
    total_off = total_spil = total_col = total_near = 0
    small_fonts = []

    for idx, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in _iter_shapes(slide.shapes):
            try:
                x, y = (sh.left or 0) / E, (sh.top or 0) / E
                w, h = (sh.width or 0) / E, (sh.height or 0) / E
            except Exception:
                continue

            # 非文本框只查越界
            if not getattr(sh, 'has_text_frame', False) or not sh.text_frame.text.strip():
                boxes.append(dict(x=x, y=y, w=w, h=h, bottom=y + h, right=x + w,
                                  need=h, over=0.0, px=0, text='', kind='shape',
                                  name=sh.name or ''))
                continue

            tf = sh.text_frame
            ml = (tf.margin_left or 0) / E
            mr = (tf.margin_right or 0) / E
            mt = (tf.margin_top or 0) / E
            mb = (tf.margin_bottom or 0) / E
            fit_w = max(10.0, w - ml - mr)
            fit_h = max(1.0, h - mt - mb)

            need = 0.0
            max_px = 0.0
            for para in tf.paragraphs:
                txt = ''.join(r.text for r in para.runs) or para.text or ''
                px = max([r.font.size.pt * k for r in para.runs if r.font.size] or [0]) or 0
                if not px:
                    continue
                max_px = max(max_px, px)
                if px < min_font:
                    small_fonts.append(dict(page=idx, px=round(px, 1), text=txt[:28]))
                family = _font_family(sh, para)
                lsp = para.line_spacing
                if lsp is not None and hasattr(lsp, 'pt'):
                    lh = float(lsp.pt) * k          # 绝对行距：直读
                else:
                    ratio = FONT_RATIO.get((family or '').strip().lower(), SINGLE_LINE_RATIO)
                    lh = px * ratio * (float(lsp) if lsp is not None else 1.05)
                need += _wrap_lines(txt, px, fit_w) * lh

            over = need - fit_h
            boxes.append(dict(x=x, y=y, w=w, h=h, bottom=y + h, right=x + w,
                              need=round(need, 1), over=round(over, 1),
                              px=round(max_px, 1), text=tf.text[:40].replace('\n', ' '),
                              kind='text', name=sh.name or ''))

        # ① 越界：x<0 视为刻意的出血装饰，不判
        off = [b for b in boxes
               if (b['bottom'] > CANVAS_H + tolerance or b['right'] > CANVAS_W + tolerance)
               and b['x'] >= -1]
        # ② 溢出：容差按字号给（中文字体末行下方本身有 leading，0.25em 内不算真溢出）
        def _thr(b):
            return max(tolerance, 0.25 * (b['px'] or 16))
        spil = [b for b in boxes
                if b['kind'] == 'text' and b['over'] > _thr(b)]
        # ②b 临界：已经超了但没超容差 —— 不算不合格，但必须报，否则隐患从缝里溜过去。
        #     下限滤掉 0.5px 级的量字噪声（PIL 度量与 PowerPoint 渲染本身有细微差异）。
        near = [b for b in boxes
                if b['kind'] == 'text'
                and max(1.0, 0.08 * (b['px'] or 16)) < b['over'] <= _thr(b)]
        # ③ 碰撞：溢出框压到下方同列形状
        cols = []
        for b in spil:
            for c in boxes:
                if c is b or c['y'] < b['y'] - 1:
                    continue
                xov = min(b['right'], c['right']) - max(b['x'], c['x'])
                intrude = (b['y'] + b['need']) - c['y']
                if xov > 20 and intrude > 6:
                    cols.append(dict(overlaps=b['text'][:26], victim=(c['text'] or c['name'])[:26],
                                     intrude_px=round(intrude, 1)))

        if off or spil or near:
            pages.append(dict(page=idx, off=off, spill=spil, collide=cols[:6], near=near))
            total_off += len(off)
            total_spil += len(spil)
            total_col += len(cols)
            total_near += len(near)

    return dict(file=os.path.abspath(pptx), slides=len(prs.slides),
                canvas=[round(CANVAS_W), round(CANVAS_H)],
                total_off=total_off, total_spill=total_spil, total_collide=total_col,
                total_near=total_near,
                pages=pages, small_fonts=small_fonts,
                pil=bool(_have_pil and _font_path), font=_font_path)


def main():
    ap = argparse.ArgumentParser(description='内容可见性体检（越界 / 溢出 / 碰撞）')
    ap.add_argument('pptx')
    ap.add_argument('--json', help='把明细写成 JSON')
    ap.add_argument('--min-font', type=float, default=0,
                    help='最小字号(px)告警阈值，0=不查')
    ap.add_argument('--tolerance', type=float, default=1.5,
                    help='几何容差(px)，默认 1.5')
    ap.add_argument('--fonts', help='量字用的字体文件路径')
    ap.add_argument('--quiet', action='store_true', help='只出总结论')
    a = ap.parse_args()

    if not os.path.exists(a.pptx):
        sys.exit('[FATAL] 文件不存在：%s' % a.pptx)

    r = audit(a.pptx, a.min_font, a.tolerance, a.fonts)

    if not r['pil']:
        print('[WARN] 缺 Pillow 或找不到字体，按字符数估算，精度下降。\n')
    elif not a.fonts:
        print('[INFO] 量字字体：%s\n' % r['font'])

    print('文件：%s' % r['file'])
    print('页数：%d ｜ 画布：%dx%d px\n' % (r['slides'], r['canvas'][0], r['canvas'][1]))

    if not a.quiet:
        hdr = '%-6s %4s %7s %5s %5s' % ('页', '越界', '溢出框', '碰撞', '临界')
        print(hdr)
        print('-' * len(hdr))
        for p in r['pages']:
            print('%-6d %4d %7d %5d %5d' % (p['page'], len(p['off']),
                                            len(p['spill']), len(p['collide']),
                                            len(p.get('near', []))))
        # 最严重的几处，直接点出来
        worst = []
        for p in r['pages']:
            for b in p['spill']:
                worst.append((b['over'], p['page'], b))
        worst.sort(key=lambda t: t[0], reverse=True)
        if worst:
            print('\n排得最狠的几处：')
            for over, pg, b in worst[:6]:
                print('  P%-3d 差 %+6.0fpx  「%s」' % (pg, -over, b['text']))
        near_all = [(p['page'], b) for p in r['pages'] for b in p.get('near', [])]
        if near_all:
            print('\n临界（超了但在容差内，还没算不合格）%d 处（前 6）：' % len(near_all))
            for pg, b in near_all[:6]:
                print('  P%-3d 超 %5.1fpx  「%s」' % (pg, b['over'], b['text']))
            print('  提示：字体替换、版本差异或对方机器缺字，都可能让这些框真的溢出。')
        if r['small_fonts']:
            print('\n小于阈值的字号 %d 处（前 6）：' % len(r['small_fonts']))
            for s in r['small_fonts'][:6]:
                print('  P%-3d %5.1fpx  「%s」' % (s['page'], s['px'], s['text']))
        print()

    ok = r['total_off'] == 0 and r['total_spill'] == 0 and r['total_collide'] == 0
    if ok:
        print('结论：合格 —— 受影响页 0 / %d，越界 0，溢出 0，碰撞 0%s' % (
            r['slides'], ('（另有临界 %d 处，未计入不合格）' % r['total_near']) if r['total_near'] else ''))
    else:
        print('结论：不合格 —— 受影响页 %d / %d，越界 %d，溢出 %d，碰撞 %d%s' % (
            len(r['pages']), r['slides'], r['total_off'], r['total_spill'],
            r['total_collide'],
            ('，另临界 %d' % r['total_near']) if r['total_near'] else ''))
        print('注意：这一层与可编辑性是两个独立判据，两边都要过才算放行。')

    if a.json:
        with open(a.json, 'w', encoding='utf-8') as f:
            json.dump(r, f, ensure_ascii=False, indent=1)
        print('明细已写入：%s' % a.json)

    return 0 if ok else (2 if r['total_collide'] else 1)


if __name__ == '__main__':
    sys.exit(main())
