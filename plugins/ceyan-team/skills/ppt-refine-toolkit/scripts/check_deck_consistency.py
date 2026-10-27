#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""跨页一致性检查：查全篇「看起来不专业」的那类问题。

单页可编辑性由 check_editability.py 负责，本脚本只管跨页：

  1. 画布尺寸是否统一
  2. 字号使用清单 —— 用了多少种、哪些只在一页出现、是否越出令牌阶梯
  3. 颜色使用清单 —— 用了多少种、哪些越出令牌、哪些只用一次
  4. 页码位置漂移 —— 缺页码、位置差超过容差
  5. 标题顶端位置 —— 同类页标题是否在同一条线上
  6. 视觉密度 —— 连续多页密度雷同（听众疲劳）、疑似过载页

用法：
    python check_deck_consistency.py <pptx路径>
    python check_deck_consistency.py <pptx路径> --tokens tokens.json
    python check_deck_consistency.py <pptx路径> --json out.json --tol 2.0

tokens.json（可选，来自 design_spec）：
    {
      "font_sizes": [12, 14, 16, 20, 24, 32, 44, 64],
      "colors": ["#1F2A2E", "#2F5D62", "#C8A46A"],
      "canvas": {"width_emu": 12192000, "height_emu": 6858000}
    }
    没有 tokens 时只做「使用频次 + 异常项」统计，不做越界判定。

退出码：0 无问题；1 存在问题；2 参数或读取错误。
"""

import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.stderr.write("缺少依赖：pip install python-pptx\n")
    sys.exit(2)

EMU_PER_PT = 12700

PAGE_NO_RE = re.compile(r"^[\-—–]?\s*\d{1,3}\s*[\-—–]?$|^\d{1,3}\s*/\s*\d{1,3}$")


# ---------------------------------------------------------------- 工具

def emu_to_pt(v):
    try:
        return round(int(v) / EMU_PER_PT, 1)
    except Exception:
        return None


def rgb_hex(color):
    """从 ColorFormat 取 #RRGGBB，取不到返回 None（主题色/继承色）"""
    try:
        if color is None:
            return None
        if color.type is None:
            return None
        rgb = color.rgb
        if rgb is None:
            return None
        return "#%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])
    except Exception:
        return None


def iter_shapes(shapes):
    """递归展开组合形状"""
    for sh in shapes:
        yield sh
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub in iter_shapes(sh.shapes):
                    yield sub
        except Exception:
            continue


# ---------------------------------------------------------------- 采集

def collect(prs):
    pages = []
    for idx, slide in enumerate(prs.slides, start=1):
        rec = {
            "page": idx,
            "sizes": [],          # [(pt, text_preview)]
            "colors": [],         # [(hex, source)]
            "inherited_runs": 0,
            "chars": 0,
            "shapes": 0,
            "page_no": None,      # {"text","x","y"}
            "title": None,        # {"size","x","y","text"}
        }
        max_size_on_page = 0.0
        max_size_shape = None
        for sh in iter_shapes(slide.shapes):
            rec["shapes"] += 1
            # 形状填充色
            try:
                if sh.fill is not None and sh.fill.type is not None:
                    h = rgb_hex(sh.fill.fore_color)
                    if h:
                        rec["colors"].append((h, "fill"))
            except Exception:
                pass
            if not getattr(sh, "has_text_frame", False):
                continue
            try:
                text = sh.text_frame.text or ""
            except Exception:
                text = ""
            rec["chars"] += len(text.strip())
            if not text.strip():
                continue
            # 字号与文字颜色
            local_max = 0.0
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        sz = r.font.size
                    except Exception:
                        sz = None
                    if sz is not None:
                        pt = round(sz.pt, 1)
                        rec["sizes"].append((pt, r.text[:12]))
                        local_max = max(local_max, pt)
                    else:
                        rec["inherited_runs"] += 1
                    try:
                        h = rgb_hex(r.font.color)
                    except Exception:
                        h = None
                    if h:
                        rec["colors"].append((h, "text"))
            if local_max > max_size_on_page:
                max_size_on_page = local_max
                max_size_shape = (sh, text.strip().replace("\n", " ")[:24])
            # 页码候选
            stripped = text.strip()
            if len(stripped) <= 8 and PAGE_NO_RE.match(stripped):
                try:
                    top = int(sh.top or 0)
                    left = int(sh.left or 0)
                except Exception:
                    top = left = 0
                if top >= int(prs.slide_height) * 0.80:
                    rec["page_no"] = {
                        "text": stripped,
                        "x": emu_to_pt(left),
                        "y": emu_to_pt(top),
                    }
        if max_size_shape:
            sh, preview = max_size_shape
            try:
                top = int(sh.top or 0)
                left = int(sh.left or 0)
            except Exception:
                top = left = 0
            rec["title"] = {
                "size": max_size_on_page,
                "x": emu_to_pt(left),
                "y": emu_to_pt(top),
                "text": preview,
            }
        pages.append(rec)
    return pages


# ---------------------------------------------------------------- 检查

def check_canvas(prs, tokens):
    w, h = int(prs.slide_width), int(prs.slide_height)
    problems = []
    want = (tokens or {}).get("canvas") or {}
    tw, th = want.get("width_emu"), want.get("height_emu")
    if tw and th and (int(tw), int(th)) != (w, h):
        problems.append(
            "画布与令牌不符：实际 %dx%d EMU，令牌 %dx%d" % (w, h, int(tw), int(th))
        )
    return {
        "width_emu": w,
        "height_emu": h,
        "width_pt": emu_to_pt(w),
        "height_pt": emu_to_pt(h),
        "ratio": round(w / h, 3) if h else None,
        "problems": problems,
    }


def check_font_sizes(pages, tokens):
    by_size = defaultdict(list)
    for p in pages:
        for pt, _ in p["sizes"]:
            by_size[pt].append(p["page"])
    out = []
    for pt in sorted(by_size, reverse=True):
        out.append({
            "pt": pt,
            "pages": sorted(set(by_size[pt])),
            "count": len(by_size[pt]),
        })
    single = [r for r in out if len(r["pages"]) == 1]
    ladder = (tokens or {}).get("font_sizes")
    off_ladder = []
    if ladder:
        lset = set(float(x) for x in ladder)
        off_ladder = [r for r in out if r["pt"] not in lset]
    problems = []
    if len(out) > 7:
        problems.append("全篇显式字号 %d 种（> 7），阶梯被稀释" % len(out))
    if single:
        problems.append(
            "只在一页出现的字号 %d 种：%s"
            % (len(single), "、".join("%gpt(P%s)" % (r["pt"], r["pages"][0]) for r in single))
        )
    if off_ladder:
        problems.append(
            "越出令牌阶梯 %d 种：%s（令牌 %s）"
            % (len(off_ladder), "、".join("%gpt" % r["pt"] for r in off_ladder),
               "/".join("%g" % float(x) for x in ladder))
        )
    return {"distribution": out, "single_page": single,
            "off_ladder": off_ladder, "problems": problems}


def check_colors(pages, tokens):
    by_color = defaultdict(list)
    src = defaultdict(Counter)
    for p in pages:
        for h, where in p["colors"]:
            by_color[h].append(p["page"])
            src[h][where] += 1
    out = []
    for h in sorted(by_color, key=lambda c: -len(by_color[c])):
        out.append({
            "hex": h,
            "pages": sorted(set(by_color[h])),
            "count": len(by_color[h]),
            "roles": dict(src[h]),
        })
    palette = (tokens or {}).get("colors")
    off = []
    if palette:
        pset = set(c.upper() for c in palette)
        off = [r for r in out if r["hex"].upper() not in pset]
    once = [r for r in out if len(r["pages"]) == 1]
    problems = []
    if off:
        problems.append(
            "令牌外颜色 %d 个：%s"
            % (len(off), "、".join("%s(P%s)" % (r["hex"], r["pages"][0]) for r in off[:8]))
        )
    if len(once) > 2 and not palette:
        problems.append("只用一次的颜色 %d 个（疑似临时取色，无令牌无法判定越界）" % len(once))
    return {"distribution": out, "off_palette": off,
            "single_use": once, "problems": problems}


def check_page_numbers(pages, tol):
    found = [p for p in pages if p["page_no"]]
    missing = [p["page"] for p in pages if not p["page_no"]]
    problems = []
    drift = []
    if found:
        mode_x = Counter(round(p["page_no"]["x"]) for p in found).most_common(1)[0][0]
        mode_y = Counter(round(p["page_no"]["y"]) for p in found).most_common(1)[0][0]
        for p in found:
            dx = round(p["page_no"]["x"] - mode_x, 1)
            dy = round(p["page_no"]["y"] - mode_y, 1)
            if abs(dx) > tol or abs(dy) > tol:
                drift.append({"page": p["page"], "dx": dx, "dy": dy,
                              "x": p["page_no"]["x"], "y": p["page_no"]["y"]})
    if missing:
        problems.append("缺页码 %d 页：%s" % (len(missing), _compress(missing)))
    if drift:
        worst = max(drift, key=lambda d: max(abs(d["dx"]), abs(d["dy"])))
        problems.append(
            "页码位置漂移 %d 页（容差 %gpt，最大 P%s dx=%g dy=%g）"
            % (len(drift), tol, worst["page"], worst["dx"], worst["dy"])
        )
    return {"detected": len(found), "missing": missing, "drift": drift,
            "anchor": {"x": mode_x if found else None, "y": mode_y if found else None},
            "problems": problems}


def check_titles(pages, tol):
    """标题顶端位置：找主簇，报离群页（待复核，不判违规）"""
    titles = [p for p in pages if p["title"]]
    if len(titles) < 4:
        return {"clusters": [], "outliers": [], "problems": []}
    buckets = Counter(round(p["title"]["y"] / 10) * 10 for p in titles)
    main_y, main_n = buckets.most_common(1)[0]
    outliers = []
    for p in titles:
        if abs(p["title"]["y"] - main_y) > tol:
            outliers.append({"page": p["page"], "y": p["title"]["y"], "dy": round(p["title"]["y"] - main_y, 1)})
    problems = []
    if len(outliers) > len(titles) * 0.4:
        problems.append(
            "标题顶端位置分散：主簇 y=%gpt 仅 %d/%d 页，%d 页偏离（可能是刻意区分，需人工判断）"
            % (main_y, main_n, len(titles), len(outliers))
        )
    return {"main_y": main_y, "main_count": main_n, "total": len(titles),
            "outliers": outliers, "problems": problems}


def check_density(pages):
    vals = [(p["page"], p["chars"], p["shapes"]) for p in pages]
    over = [{"page": pg, "chars": c} for pg, c, _ in vals if c >= 380]
    runs, cur = [], []
    for pg, c, _ in vals:
        if cur and abs(c - cur[-1][1]) <= 10 and c >= 40:
            cur.append((pg, c))
        else:
            if len(cur) >= 5:
                runs.append(cur)
            cur = [(pg, c)]
    if len(cur) >= 5:
        runs.append(cur)
    problems = []
    if over:
        problems.append("疑似过载 %d 页（正文字符 ≥ 380）：%s"
                        % (len(over), "、".join("P%02d(%d字)" % (o["page"], o["chars"]) for o in over[:6])))
    if runs:
        problems.append(
            "连续 %d 段密度雷同（≥5 页字数几乎相同）：%s"
            % (len(runs), "；".join("P%02d–P%02d" % (r[0][0], r[-1][0]) for r in runs))
        )
    return {"overload": over,
            "flat_runs": [{"from": r[0][0], "to": r[-1][0], "chars": r[0][1]} for r in runs],
            "problems": problems}


def _compress(nums):
    nums = sorted(nums)
    out, start, prev = [], None, None
    for n in nums:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append("P%02d" % start if start == prev else "P%02d–P%02d" % (start, prev))
            start = prev = n
    if start is not None:
        out.append("P%02d" % start if start == prev else "P%02d–P%02d" % (start, prev))
    return "、".join(out)


# ---------------------------------------------------------------- 主流程

def main():
    ap = argparse.ArgumentParser(description="pptx 跨页一致性检查")
    ap.add_argument("pptx")
    ap.add_argument("--tokens", help="design_spec 令牌 json，用于越界判定")
    ap.add_argument("--json", dest="json_out", help="明细输出路径")
    ap.add_argument("--tol", type=float, default=2.0, help="页码位置容差(pt)，默认 2")
    args = ap.parse_args()

    if not os.path.isfile(args.pptx):
        sys.stderr.write("文件不存在：%s\n" % args.pptx)
        return 2
    tokens = None
    if args.tokens:
        try:
            with open(args.tokens, encoding="utf-8") as f:
                tokens = json.load(f)
        except Exception as e:
            sys.stderr.write("tokens 读取失败：%s\n" % e)
            return 2
    try:
        prs = Presentation(args.pptx)
    except Exception as e:
        sys.stderr.write("pptx 读取失败：%s\n" % e)
        return 2

    pages = collect(prs)
    canvas = check_canvas(prs, tokens)
    fonts = check_font_sizes(pages, tokens)
    colors = check_colors(pages, tokens)
    pnos = check_page_numbers(pages, args.tol)
    titles = check_titles(pages, 6.0)
    dens = check_density(pages)

    sections = [
        ("画布", canvas),
        ("字号使用", fonts),
        ("颜色使用", colors),
        ("页码位置", pnos),
        ("标题顶端位置", titles),
        ("视觉密度", dens),
    ]

    W = 78
    print("=" * W)
    print("跨页一致性检查：%s（%d 页）" % (os.path.basename(args.pptx), len(pages)))
    print("=" * W)

    print("\n[画布]")
    print("  %dx%d EMU  =  %gpt x %gpt  比例 %g"
          % (canvas["width_emu"], canvas["height_emu"],
             canvas["width_pt"], canvas["height_pt"], canvas["ratio"]))

    print("\n[字号使用] 显式设定 %d 种" % len(fonts["distribution"]))
    for r in fonts["distribution"]:
        mark = " ← 仅一页" if len(r["pages"]) == 1 else ""
        print("  %-7gpt %3d 处  %s%s"
              % (r["pt"], r["count"], _compress(r["pages"])[:52], mark))
    inh = sum(p["inherited_runs"] for p in pages)
    if inh:
        print("  （继承字号 %d 处，未计入上表——它们来自版式/母版）" % inh)

    print("\n[颜色使用] %d 种" % len(colors["distribution"]))
    for r in colors["distribution"][:12]:
        roles = "/".join("%s%d" % (k, v) for k, v in r["roles"].items())
        print("  %-9s %3d 处  %-6s %s"
              % (r["hex"], r["count"], roles, _compress(r["pages"])[:46]))
    if len(colors["distribution"]) > 12:
        print("  …… 另 %d 种" % (len(colors["distribution"]) - 12))

    print("\n[页码位置] 检出 %d/%d 页" % (pnos["detected"], len(pages)))
    if pnos["drift"]:
        for d in pnos["drift"][:8]:
            print("  P%02d  x=%g y=%g  (dx=%+g dy=%+g)"
                  % (d["page"], d["x"], d["y"], d["dx"], d["dy"]))

    print("\n[标题顶端位置] 主簇 y=%gpt（%d/%d 页）"
          % (titles.get("main_y", 0), titles.get("main_count", 0), titles.get("total", 0)))
    if titles["outliers"]:
        print("  偏离：%s" % "、".join("P%02d(%+gpt)" % (o["page"], o["dy"]) for o in titles["outliers"][:10]))

    print("\n[视觉密度]")
    print("  字数 %d–%d，形状 %d–%d"
          % (min(p["chars"] for p in pages), max(p["chars"] for p in pages),
             min(p["shapes"] for p in pages), max(p["shapes"] for p in pages)))

    all_problems = []
    for name, sec in sections:
        for msg in sec.get("problems", []):
            all_problems.append((name, msg))

    print("\n" + "-" * W)
    if all_problems:
        print("发现 %d 类问题：" % len(all_problems))
        for name, msg in all_problems:
            print("  ⚠ [%s] %s" % (name, msg))
    else:
        print("✅ 未发现跨页一致性问题")
    print("-" * W)

    detail = {
        "file": os.path.abspath(args.pptx),
        "pages": len(pages),
        "canvas": canvas,
        "font_sizes": fonts,
        "colors": colors,
        "page_numbers": pnos,
        "titles": titles,
        "density": dens,
        "problems": [{"section": n, "message": m} for n, m in all_problems],
        "per_page": [
            {"page": p["page"], "chars": p["chars"], "shapes": p["shapes"],
             "page_no": p["page_no"], "title": p["title"]}
            for p in pages
        ],
    }
    out = args.json_out or os.path.splitext(args.pptx)[0] + ".consistency.json"
    try:
        with open(out, "w", encoding="utf-8") as f:
            json.dump(detail, f, ensure_ascii=False, indent=2)
        print("明细已写入：%s" % out)
    except Exception as e:
        sys.stderr.write("JSON 写入失败：%s\n" % e)

    return 1 if all_problems else 0


if __name__ == "__main__":
    sys.exit(main())
